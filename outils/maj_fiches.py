# -*- coding: utf-8 -*-
"""
Mise à jour de fiches.json quand tu ajoutes un PDF, un .docx ou une photo.

Ce que fait le script :
  1. Il cherche les PDF / .docx / photos (.jpg, .jpeg, .png, .heic, .webp) :
       - dans le dossier principal « exercice golf »  -> exercices de GOLF
       - dans le sous-dossier « prepa phisique »      -> exercices PHYSIQUES
  2. Pour chaque fichier qui n'a encore AUCUNE fiche dans fiches.json,
     il crée une fiche « brouillon » (avec le texte extrait pour un PDF/.docx).
     Une photo ne contient pas de texte lisible par le script : le brouillon
     dit simplement « photo à analyser » (une photo peut contenir plusieurs
     exercices, c'est Claude qui les rédige).
     Un brouillon n'apparaît PAS dans le programme du jour : il est visible
     dans l'onglet « Bibliothèque » avec la mention « À compléter ».
     Les fichiers listés dans « fichiersAnalyses » (doublons, tableaux…)
     sont considérés comme déjà traités.
  3. Il signale les PDF illisibles (aucun texte : souvent un PDF scanné).
  4. Il signale les fiches dont le fichier source a disparu (sans les supprimer).
  5. Il régénère fiches-data.js (copie de secours de fiches.json, utilisée
     quand on ouvre index.html en double-cliquant, sans serveur).

Le script ne modifie JAMAIS tes PDF, .docx ou Excel : il les lit seulement.
Avant d'écrire, il garde une copie de l'ancien fichier : fiches.json.bak

Lancement : double-clic sur « maj_fiches.bat » (dans le même dossier).
"""

import json
import re
import shutil
import sys
import unicodedata
from datetime import date
from pathlib import Path

# --- Emplacements ---------------------------------------------------------
DOSSIER_OUTILS = Path(__file__).resolve().parent
DOSSIER_SITE = DOSSIER_OUTILS.parent                    # « exercice golf »
DOSSIER_PHYSIQUE = DOSSIER_SITE / "prepa phisique"
FICHIER_JSON = DOSSIER_SITE / "fiches.json"
FICHIER_JS = DOSSIER_SITE / "fiches-data.js"

# Champs obligatoires d'une fiche « valide » (vérifiés à chaque lancement)
CHAMPS_OBLIGATOIRES = [
    "id", "titre", "type", "categorie", "duree", "objectif", "materiel",
    "miseEnPlace", "deroule", "reussite", "varianteFacile",
    "varianteDifficile", "source",
]
EXTENSIONS_DOCUMENTS = (".pdf", ".docx")
EXTENSIONS_PHOTOS = (".jpg", ".jpeg", ".png", ".heic", ".webp")
CATEGORIES = {
    "golf": ["putting", "chipping", "approches", "fers", "drive", "bunker"],
    "physique": ["cardio", "force", "puissance", "mobilite", "abdos", "recuperation"],
}


def lire_texte_pdf(chemin):
    """Renvoie le texte d'un PDF, ou None s'il est illisible."""
    from pypdf import PdfReader
    try:
        lecteur = PdfReader(str(chemin))
        pages = [(page.extract_text() or "").strip() for page in lecteur.pages]
        return "\n\n".join(p for p in pages if p)
    except Exception:
        return None


def lire_texte_docx(chemin):
    """Renvoie le texte d'un .docx (paragraphes + tableaux), ou None."""
    import docx
    try:
        document = docx.Document(str(chemin))
        lignes = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for tableau in document.tables:
            for rangee in tableau.rows:
                lignes.append(" | ".join(c.text.strip() for c in rangee.cells))
        return "\n".join(lignes)
    except Exception:
        return None


def titre_depuis_nom(nom_fichier):
    """« Demi Lune.pdf.pdf » -> « Demi Lune »."""
    titre = nom_fichier
    while titre.lower().endswith(EXTENSIONS_DOCUMENTS + EXTENSIONS_PHOTOS):
        titre = titre.rsplit(".", 1)[0]
    return titre.strip()


def identifiant(texte):
    """« Équilibre unipodal » -> « equilibre-unipodal »."""
    sans_accent = unicodedata.normalize("NFKD", texte).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", sans_accent.lower()).strip("-") or "fiche"


def lister_sources():
    """Liste (chemin relatif, type) des documents à prendre en compte."""
    sources = []
    extensions = EXTENSIONS_DOCUMENTS + EXTENSIONS_PHOTOS
    for chemin in sorted(DOSSIER_SITE.iterdir()):
        if chemin.is_file() and chemin.suffix.lower() in extensions:
            sources.append((chemin.name, "golf"))
    if DOSSIER_PHYSIQUE.is_dir():
        for chemin in sorted(DOSSIER_PHYSIQUE.rglob("*")):
            if chemin.is_file() and chemin.suffix.lower() in extensions:
                relatif = chemin.relative_to(DOSSIER_SITE).as_posix()
                sources.append((relatif, "physique"))
    # Les fichiers temporaires de Word (« ~$... ») sont ignorés
    return [(r, t) for (r, t) in sources if not Path(r).name.startswith("~$")]


def nfc(texte):
    """Forme Unicode standard : un « É » peut être stocké de 2 façons différentes
    selon le logiciel qui a créé le fichier. On compare toujours en forme NFC."""
    return unicodedata.normalize("NFC", texte or "")


def verifier_fiche(fiche):
    """Renvoie la liste des problèmes d'une fiche valide (vide si tout va bien)."""
    problemes = [f"champ « {c} » manquant" for c in CHAMPS_OBLIGATOIRES if not fiche.get(c)]
    type_fiche = fiche.get("type")
    if type_fiche not in CATEGORIES:
        problemes.append(f"type « {type_fiche} » inconnu (golf ou physique)")
    elif fiche.get("categorie") not in CATEGORIES[type_fiche]:
        problemes.append(f"catégorie « {fiche.get('categorie')} » inconnue pour le type {type_fiche}")
    return problemes


def main():
    print("=== Mise à jour des fiches ===\n")

    donnees = json.loads(FICHIER_JSON.read_text(encoding="utf-8"))
    fiches = donnees["fiches"]

    sources = lister_sources()
    # Nom « normalisé » -> nom exact sur le disque
    noms_reels = {nfc(relatif): relatif for relatif, _ in sources}

    # Chaque fiche reprend le nom exact du disque (indispensable pour que le
    # lien vers le PDF fonctionne une fois en ligne)
    for f in fiches:
        source = f.get("source") or {}
        reel = noms_reels.get(nfc(source.get("fichier")))
        if reel and reel != source["fichier"]:
            source["fichier"] = reel
        if source.get("fichiers"):
            source["fichiers"] = [noms_reels.get(nfc(x), x) for x in source["fichiers"]]

    # Fichiers déjà couverts : par une fiche (fichier principal ou liste
    # « fichiers » quand un exercice apparaît sur plusieurs photos), ou
    # marqués comme déjà analysés (doublons, tableaux sans exercice…)
    deja_couverts = set()
    for f in fiches:
        source = f.get("source") or {}
        deja_couverts.add(nfc(source.get("fichier")))
        deja_couverts.update(nfc(x) for x in source.get("fichiers", []))
    deja_couverts.update(nfc(x) for x in donnees.get("fichiersAnalyses", {}))
    ids_existants = {f["id"] for f in fiches}

    nouvelles = 0
    for relatif, type_fiche in sources:
        if nfc(relatif) in deja_couverts:
            continue

        chemin = DOSSIER_SITE / relatif
        suffixe = chemin.suffix.lower()
        if suffixe == ".pdf":
            texte = lire_texte_pdf(chemin)
            type_source = "pdf"
        elif suffixe == ".docx":
            texte = lire_texte_docx(chemin)
            type_source = "docx"
        else:
            texte = None
            type_source = "photo"

        titre = titre_depuis_nom(chemin.name)
        nouvel_id = identifiant(titre)
        while nouvel_id in ids_existants:          # évite les doublons d'identifiant
            nouvel_id += "-bis"
        ids_existants.add(nouvel_id)

        fiche = {
            "id": nouvel_id,
            "titre": titre,
            "type": type_fiche,
            "categorie": "a-classer",
            "duree": 15,
            "objectif": "",
            "materiel": [],
            "miseEnPlace": [],
            "deroule": [],
            "reussite": "",
            "varianteFacile": "",
            "varianteDifficile": "",
            "source": {"type": type_source, "fichier": relatif},
            "statut": "brouillon",
            "texteExtrait": (texte or "")[:3000],
        }
        if type_source == "photo":
            fiche["alerte"] = ("Nouvelle photo à analyser : demande à Claude de la lire "
                               "et de rédiger les fiches des exercices qu'elle contient.")
            print(f"  + Nouvelle photo à analyser : {relatif}")
        elif not texte:
            fiche["alerte"] = ("Aucun texte lisible dans ce fichier (document scanné ou image ?). "
                               "Il faudra décrire l'exercice à la main.")
            print(f"  ⚠ ILLISIBLE : {relatif} -> brouillon créé SANS texte")
        else:
            print(f"  + Nouveau brouillon : {relatif}  (id : {nouvel_id})")
        fiches.append(fiche)
        nouvelles += 1

    # Fiches dont le fichier source n'existe plus (on prévient seulement)
    for f in fiches:
        source = f.get("source") or {}
        for fichier in [source.get("fichier")] + source.get("fichiers", []):
            if fichier and not (DOSSIER_SITE / fichier).exists():
                print(f"  ⚠ Fichier introuvable pour la fiche « {f['titre']} » : {fichier}")

    # Vérification des fiches validées
    for f in fiches:
        if f.get("statut") == "valide":
            for probleme in verifier_fiche(f):
                print(f"  ⚠ Fiche « {f.get('titre', '?')} » : {probleme}")

    brouillons = [f for f in fiches if f.get("statut") == "brouillon"]
    donnees["miseAJour"] = date.today().isoformat()

    # Sauvegarde de l'ancienne version, puis écriture
    shutil.copyfile(FICHIER_JSON, FICHIER_JSON.with_suffix(".json.bak"))
    texte_json = json.dumps(donnees, ensure_ascii=False, indent=2)
    FICHIER_JSON.write_text(texte_json + "\n", encoding="utf-8")
    FICHIER_JS.write_text(
        "// Fichier généré automatiquement par outils/maj_fiches.py : ne pas modifier à la main.\n"
        "// Copie de fiches.json, utilisée quand la page est ouverte sans serveur (double-clic).\n"
        "window.FICHES_SECOURS = " + texte_json + ";\n",
        encoding="utf-8",
    )

    valides = len(fiches) - len(brouillons)
    print(f"\n{nouvelles} nouvelle(s) fiche(s) brouillon.")
    print(f"Total : {len(fiches)} fiches ({valides} validées, {len(brouillons)} à compléter).")
    if brouillons:
        print("\nÀ faire : demande à Claude de rédiger les brouillons suivants :")
        for f in brouillons:
            print(f"   - {f['titre']}  ({f['source']['fichier']})")
    print("\nTerminé. Pense à publier les changements sur GitHub (voir le guide).")


if __name__ == "__main__":
    try:
        main()
    except ModuleNotFoundError as erreur:
        print(f"Module Python manquant : {erreur.name}.")
        print("Installe-le avec : python -m pip install pypdf python-docx")
        sys.exit(1)
    except json.JSONDecodeError as erreur:
        print("fiches.json n'est pas un JSON valide (virgule ou guillemet en trop ?).")
        print(f"Ligne {erreur.lineno}, colonne {erreur.colno}. Restaure fiches.json.bak si besoin.")
        sys.exit(1)
